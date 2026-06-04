#!/usr/bin/env python3
"""
streamlit_app.py
================
Web app version of the Polymer Injectivity Calculator — eppok EURL.

Run locally:
    streamlit run streamlit_app.py

Deploy free on Streamlit Community Cloud (see DEPLOY_WEB.md).

Password protection:
    Set a password in .streamlit/secrets.toml (local) or in the
    Streamlit Cloud "Secrets" panel:

        app_password = "your-password-here"

    If no password is configured, the app runs open (useful for local dev).
"""

import io
import hmac
from pathlib import Path

import numpy as np
import pandas as pd
import streamlit as st
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from polymer_injectivity.engine import (
    make_model,
    normalize_config,
    load_rate_schedule_from_excel,
    rate_schedule_summary,
    BASE_DEFAULTS,
    MODEL_DEFAULTS,
)
from polymer_injectivity.presets import (
    PARAM_GROUPS,
    INT_PARAMS,
    PRESETS,
    MODEL_LABELS,
)

# ──────────────────────────────────────────────────────────────────────────────
#  PAGE CONFIG + THEME
# ──────────────────────────────────────────────────────────────────────────────

st.set_page_config(
    page_title="Polymer Injectivity Calculator — eppok",
    page_icon="🛢️",
    layout="wide",
    initial_sidebar_state="expanded",
)

NAVY = "#1F3A5F"
NAVY_DARK = "#13243B"
ORANGE = "#E87722"
LIGHT = "#F5F7FA"

st.markdown(f"""
<style>
    .stApp {{ background: {LIGHT}; }}
    /* Header band */
    .eppok-header {{
        background: linear-gradient(100deg, {NAVY_DARK} 0%, {NAVY} 70%, {NAVY} 100%);
        padding: 1.4rem 1.8rem; border-radius: 12px; margin-bottom: 1.2rem;
        border-left: 6px solid {ORANGE};
    }}
    .eppok-header h1 {{
        color: #FFFFFF; font-size: 1.55rem; margin: 0; font-weight: 700;
        letter-spacing: -0.01em;
    }}
    .eppok-header p {{
        color: #B9C6D6; font-size: 0.92rem; margin: 0.25rem 0 0 0;
    }}
    .eppok-header .accent {{ color: {ORANGE}; }}
    /* Metric cards */
    div[data-testid="stMetric"] {{
        background: #FFFFFF; border: 1px solid #E2E8F0; border-radius: 10px;
        padding: 0.9rem 1.1rem; box-shadow: 0 1px 3px rgba(19,36,59,0.06);
    }}
    div[data-testid="stMetriclabel"] p {{ color: {NAVY}; font-weight: 600; }}
    /* Run button */
    .stButton > button[kind="primary"] {{
        background: {ORANGE}; border: none; font-weight: 700; font-size: 1.02rem;
        padding: 0.55rem 0; border-radius: 8px;
    }}
    .stButton > button[kind="primary"]:hover {{ background: #C9631A; }}
    /* Sidebar */
    section[data-testid="stSidebar"] {{ background: #FFFFFF; border-right: 1px solid #E2E8F0; }}
    section[data-testid="stSidebar"] h2 {{ color: {NAVY}; font-size: 1.05rem; }}
    /* Expander headers */
    details summary {{ font-weight: 600; color: {NAVY}; }}
    /* Tighten number inputs */
    div[data-testid="stNumberInput"] label p {{ font-size: 0.82rem; }}
    .footer-note {{
        color: #8595A8; font-size: 0.8rem; text-align: center;
        margin-top: 2rem; padding-top: 1rem; border-top: 1px solid #E2E8F0;
    }}
</style>
""", unsafe_allow_html=True)


# ──────────────────────────────────────────────────────────────────────────────
#  PASSWORD GATE
# ──────────────────────────────────────────────────────────────────────────────

def check_password() -> bool:
    """Return True if the user has entered the correct password (or none is set)."""
    # If no password is configured anywhere, run open.
    configured = None
    try:
        configured = st.secrets.get("app_password", None)
    except Exception:
        configured = None

    if not configured:
        return True  # open mode (local dev / no secret set)

    if st.session_state.get("password_ok", False):
        return True

    # Centered login card
    _, mid, _ = st.columns([1, 1.3, 1])
    with mid:
        st.markdown(f"""
        <div class="eppok-header" style="text-align:center; margin-top:3rem;">
            <h1>🛢️ Polymer Injectivity Calculator</h1>
            <p><span class="accent">eppok EURL</span> · secure access</p>
        </div>
        """, unsafe_allow_html=True)
        pw = st.text_input("Password", type="password", placeholder="Enter access password")
        if st.button("Enter", type="primary", use_container_width=True):
            if hmac.compare_digest(pw, str(configured)):
                st.session_state["password_ok"] = True
                st.rerun()
            else:
                st.error("Incorrect password.")
        st.caption("Contact eppok EURL for access credentials.")
    return False


if not check_password():
    st.stop()


# ──────────────────────────────────────────────────────────────────────────────
#  HEADER
# ──────────────────────────────────────────────────────────────────────────────

st.markdown(f"""
<div class="eppok-header">
    <h1>🛢️ Polymer Injectivity Calculator</h1>
    <p><span class="accent">eppok EURL</span> · Analytical toolkit for polymer flooding injectivity
    · Delamaide · UVIM · UVIM+Fracture · Aitkulov</p>
</div>
""", unsafe_allow_html=True)


# ──────────────────────────────────────────────────────────────────────────────
#  SESSION STATE — parameter values
# ──────────────────────────────────────────────────────────────────────────────

def _default_value(key):
    """Best-effort default for a parameter key."""
    if key in BASE_DEFAULTS and BASE_DEFAULTS[key] is not None:
        v = BASE_DEFAULTS[key]
        if isinstance(v, (int, float)):
            return float(v)
    return 0.0

if "params" not in st.session_state:
    st.session_state.params = {}
    for _, plist in PARAM_GROUPS:
        for _label, key, _unit, _tip in plist:
            st.session_state.params[key] = _default_value(key)

if "excel_schedule" not in st.session_state:
    st.session_state.excel_schedule = None


def apply_preset(preset_name: str):
    """Load a preset's values into session state."""
    preset = PRESETS[preset_name]
    for _, plist in PARAM_GROUPS:
        for _label, key, _unit, _tip in plist:
            if key in preset:
                st.session_state.params[key] = float(preset[key])
    st.session_state.preset_model = preset.get("model", "vertical_uvim_fracture")
    st.session_state.preset_frac_model = preset.get("fracture_model", "PKN")
    st.session_state.preset_frac_on = preset.get("fracture_on", True)
    st.session_state.excel_schedule = None


# ──────────────────────────────────────────────────────────────────────────────
#  SIDEBAR — model, preset, options
# ──────────────────────────────────────────────────────────────────────────────

with st.sidebar:
    st.markdown("## ⚙️ Configuration")

    # Preset
    preset_choice = st.selectbox(
        "📋 Load a field preset",
        ["— none —"] + list(PRESETS.keys()),
        help="Pre-filled parameters from published field cases.",
    )
    if st.button("Apply preset", use_container_width=True) and preset_choice != "— none —":
        apply_preset(preset_choice)
        st.success(f"Loaded: {preset_choice}")

    st.divider()

    # Model
    model_keys = list(MODEL_LABELS.keys())
    default_model_idx = model_keys.index(
        st.session_state.get("preset_model", "vertical_uvim_fracture")
    ) if st.session_state.get("preset_model") in model_keys else 2
    model = st.selectbox(
        "🧮 Model",
        model_keys,
        index=default_model_idx,
        format_func=lambda k: MODEL_LABELS[k],
    )

    st.divider()

    # Options
    st.markdown("### Options")
    frac_on = st.checkbox("Enable fracture", value=st.session_state.get("preset_frac_on", True))
    frac_model = st.radio("Fracture model", ["PKN", "KGD"],
                          index=0 if st.session_state.get("preset_frac_model", "PKN") == "PKN" else 1,
                          horizontal=True)
    frac_growth = st.radio("Growth mode", ["criterion", "bhp_target"], horizontal=True)
    frac_persist = st.checkbox("Fracture persistence (monotonic growth)", value=True)
    tertiary = st.checkbox("Tertiary flood (post-waterflood Sw)", value=True,
                           help="Uncheck for secondary flood (polymer from day one).")
    use_krw_wf = st.checkbox("Use krw in waterflood zone", value=True)

    st.divider()
    st.caption("eppok EURL · polymer injectivity v1.0")


# ──────────────────────────────────────────────────────────────────────────────
#  MAIN — parameter inputs in tabs
# ──────────────────────────────────────────────────────────────────────────────

# Relevant parameter groups depend on selected model
RELEVANT_GROUPS = {
    "vertical_delamaide":     {"Reservoir", "Polymer (Power-Law)", "Injection & Time", "Densities & Other"},
    "vertical_uvim":          {"Reservoir", "Polymer (Power-Law)", "UVIM (Viscoelastic)", "Injection & Time", "Densities & Other"},
    "vertical_uvim_fracture": {"Reservoir", "Polymer (Power-Law)", "UVIM (Viscoelastic)", "Injection & Time", "Fracture (SPE-215083)", "Densities & Other"},
    "horizontal_aitkulov":    {"Reservoir", "Polymer (Power-Law)", "Injection & Time", "Horizontal Well", "Densities & Other"},
}
active_groups = RELEVANT_GROUPS.get(model, {g for g, _ in PARAM_GROUPS})

st.markdown("### 📊 Parameters")
group_names = [g for g, _ in PARAM_GROUPS if g in active_groups]
tabs = st.tabs(group_names)

for tab, gname in zip(tabs, group_names):
    with tab:
        plist = next(pl for g, pl in PARAM_GROUPS if g == gname)
        cols = st.columns(3)
        for i, (label, key, unit, tip) in enumerate(plist):
            with cols[i % 3]:
                unit_label = f"{label}" + (f"  ({unit})" if unit != "—" else "")
                is_int = key in INT_PARAMS
                val = st.session_state.params.get(key, _default_value(key))
                new_val = st.number_input(
                    unit_label,
                    value=float(val),
                    key=f"input_{key}",
                    help=tip,
                    format="%.4g",
                )
                st.session_state.params[key] = int(new_val) if is_int else new_val


# ──────────────────────────────────────────────────────────────────────────────
#  RATE SCHEDULE
# ──────────────────────────────────────────────────────────────────────────────

st.markdown("### 📅 Injection rate schedule (optional)")
rs_col1, rs_col2 = st.columns([1, 1])

with rs_col1:
    st.markdown("**Option A — Upload Excel**")
    xl_c1, xl_c2 = st.columns(2)
    with xl_c1:
        date_col = st.text_input("Date column name", value="date")
    with xl_c2:
        rate_col = st.text_input("Rate column name", value="q_bpd")

    uploaded = st.file_uploader(
        "Upload .xlsx with date + rate columns",
        type=["xlsx", "xls"],
        help="Two columns: a date and an injection rate (bbl/day). Rate is held "
             "constant until the next date.",
    )
    if uploaded is not None:
        try:
            schedule = load_rate_schedule_from_excel(
                uploaded, date_col=date_col.strip() or "date",
                rate_col=rate_col.strip() or "q_bpd",
            )
            st.session_state.excel_schedule = schedule
            st.success(f"✅ {rate_schedule_summary(schedule)}")
        except Exception as e:
            st.session_state.excel_schedule = None
            st.error(f"Could not load file: {e}")

with rs_col2:
    st.markdown("**Option B — Type manually**")
    manual_text = st.text_area(
        "One `day, rate` pair per line",
        placeholder="0, 800\n30, 1000\n90, 1200",
        height=150,
    )

# Build a downloadable template (separate, always available)
def build_template_bytes() -> bytes:
    import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Rate Schedule"
    headers = ["date", "q_bpd", "notes"]
    fill = PatternFill("solid", fgColor="1F3A5F")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 30
    examples = [
        (datetime.date(2024, 1, 1), 800, "Injection start"),
        (datetime.date(2024, 2, 1), 1000, "Step increase"),
        (datetime.date(2024, 4, 1), 1200, "Full rate"),
        (datetime.date(2024, 7, 1), 1500, "Plateau"),
    ]
    for r, (d, q, n) in enumerate(examples, start=2):
        ws.cell(row=r, column=1, value=d).number_format = "YYYY-MM-DD"
        ws.cell(row=r, column=2, value=q)
        ws.cell(row=r, column=3, value=n)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

st.download_button(
    "⬇️ Download Excel template",
    data=build_template_bytes(),
    file_name="rate_schedule_template.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)


# ──────────────────────────────────────────────────────────────────────────────
#  RUN
# ──────────────────────────────────────────────────────────────────────────────

st.markdown("")
run = st.button("▶  RUN SIMULATION", type="primary", use_container_width=True)


def assemble_config() -> dict:
    cfg = dict(st.session_state.params)
    cfg["model"] = model
    cfg["fracture_on"] = frac_on
    cfg["fracture_model"] = frac_model
    cfg["fracture_growth_mode"] = frac_growth
    cfg["fracture_persistence"] = frac_persist
    cfg["tertiary_flood"] = tertiary
    cfg["use_krw_waterflood"] = use_krw_wf

    # Rate schedule: Excel wins over manual text
    if st.session_state.excel_schedule:
        cfg["rate_schedule"] = st.session_state.excel_schedule
    elif manual_text.strip():
        sched = []
        for line in manual_text.strip().splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = line.replace(";", ",").split(",")
            if len(parts) >= 2:
                sched.append((float(parts[0]), float(parts[1])))
        if sched:
            cfg["rate_schedule"] = sched
    return cfg


def make_figure(rows):
    days = [r["day"] for r in rows]
    bhp = [r["bhp_psi"] for r in rows]
    rp = [r.get("rp_ft", 0.0) for r in rows]
    ii = [r.get("injectivity_bpd_per_psi", 0.0) for r in rows]
    Lf = [r.get("Lf_ft", 0.0) for r in rows]

    fig, axes = plt.subplots(2, 2, figsize=(11, 6.5), facecolor="white")
    ax = axes.ravel()
    plot_specs = [
        (bhp, "Bottom-hole pressure", "psi", NAVY),
        (rp, "Polymer front extent", "ft", ORANGE),
        (ii, "Injectivity index", "bpd/psi", "#2E7D5B"),
        (Lf, "Fracture half-length", "ft", "#9C2B2B"),
    ]
    for a, (series, title, ylab, color) in zip(ax, plot_specs):
        a.plot(days, series, color=color, linewidth=2)
        a.set_title(title, fontsize=11, fontweight="bold", color=NAVY)
        a.set_xlabel("day", fontsize=9)
        a.set_ylabel(ylab, fontsize=9)
        a.grid(True, alpha=0.25)
        a.set_facecolor("#FCFCFD")
    fig.tight_layout()
    return fig


if run:
    try:
        cfg = assemble_config()
        with st.spinner("Running simulation..."):
            model_obj = make_model(cfg)
            rows = model_obj.run()
        st.session_state.results = rows
        st.session_state.results_cfg = cfg
    except Exception as e:
        st.error(f"Simulation error: {e}")
        st.session_state.results = None


# ──────────────────────────────────────────────────────────────────────────────
#  RESULTS
# ──────────────────────────────────────────────────────────────────────────────

if st.session_state.get("results"):
    rows = st.session_state.results
    first, last = rows[0], rows[-1]

    st.markdown("### 📈 Results")

    # Key metrics
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Final BHP", f"{last['bhp_psi']:,.0f} psi",
              f"{last['bhp_psi'] - first['bhp_psi']:+,.0f} psi")
    m2.metric("Final injectivity", f"{last['injectivity_bpd_per_psi']:.3f} bpd/psi")
    m3.metric("Polymer front", f"{last.get('rp_ft', 0):,.1f} ft")
    max_lf = max(r.get("Lf_ft", 0.0) for r in rows)
    m4.metric("Max fracture Lf", f"{max_lf:,.1f} ft" if max_lf > 0 else "— no frac")

    # Plots
    fig = make_figure(rows)
    st.pyplot(fig, use_container_width=True)

    # Data table + downloads
    df = pd.DataFrame(rows)
    with st.expander("📋 Full data table"):
        st.dataframe(df, use_container_width=True, height=320)

    dl1, dl2 = st.columns(2)
    with dl1:
        csv_bytes = df.to_csv(index=False).encode("utf-8")
        st.download_button("⬇️ Download CSV", data=csv_bytes,
                           file_name="polymer_injectivity_results.csv",
                           mime="text/csv", use_container_width=True)
    with dl2:
        png_buf = io.BytesIO()
        fig.savefig(png_buf, format="png", dpi=150, bbox_inches="tight")
        st.download_button("⬇️ Download plot (PNG)", data=png_buf.getvalue(),
                           file_name="polymer_injectivity_plot.png",
                           mime="image/png", use_container_width=True)
    plt.close(fig)
else:
    st.info("Set your parameters above (or load a preset), then click **Run simulation**.")


st.markdown(
    '<div class="footer-note">eppok EURL · Polymer injectivity analytical toolkit · '
    'Models: SPE-195513, SPE-215083, Abdullah et al. 2023, Aitkulov et al. 2021 · '
    'For internal and client use</div>',
    unsafe_allow_html=True,
)
